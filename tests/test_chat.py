# -*- coding: utf-8 -*-
"""智能对话模块专项测试：附件解析 + 上传/会话边界行为。

直接运行：python tests/test_chat.py
- 进程内 TestClient，独立临时 SQLite / 上传目录，不污染开发库
- MOCK_LLM=1 强制 mock；UPLOAD_MAX_MB=1 以便快速触发 413
"""
from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

_TMP = Path(tempfile.mkdtemp(prefix="vflo_chat_test_"))
os.environ["MOCK_LLM"] = "1"
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/test.db"
os.environ["UPLOAD_DIR"] = str(_TMP / "uploads")
os.environ["UPLOAD_MAX_MB"] = "1"  # 1MB 上限，便于构造 413

from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402
from app.services.file_extract import extract_text, is_allowed, safe_ext  # noqa: E402

FAILED: list[str] = []


def check(name: str, cond: bool, detail: object = "") -> None:
    if cond:
        print(f"  ok  {name}")
    else:
        FAILED.append(name)
        print(f" FAIL {name} :: {detail}")


def test_extract_units() -> None:
    print("[unit] file_extract")
    d = _TMP / "extract"
    d.mkdir(parents=True, exist_ok=True)

    check("safe_ext 大小写/无后缀", safe_ext("A.PDF") == "pdf" and safe_ext("noext") == "")
    check("is_allowed 白名单", is_allowed("docx") and not is_allowed("exe"))

    p = d / "u8.txt"
    p.write_bytes("你好，世界".encode("utf-8"))
    text, status, _ = extract_text(p, "txt")
    check("txt utf-8", status == "ok" and "你好" in text, (status, text[:20]))

    p.write_bytes("国标编码".encode("gbk"))
    text, status, _ = extract_text(p, "txt")
    check("txt gbk 回退", status == "ok" and "国标" in text, (status, text[:20]))

    p.write_bytes(b"")
    text, status, note = extract_text(p, "txt")
    check("空文件 failed", status == "failed" and note, (status, note))

    try:
        import docx
        doc = docx.Document()
        doc.add_paragraph("段落一")
        t = doc.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "左"
        t.cell(0, 1).text = "右"
        p2 = d / "a.docx"
        doc.save(str(p2))
        text, status, _ = extract_text(p2, "docx")
        check("docx 段落+表格", status == "ok" and "段落一" in text and "左 | 右" in text, (status, text[:40]))
    except ImportError:
        print(" skip docx（缺少 python-docx）")

    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "进度"
        ws.append(["任务", "工时"])
        ws.append(["写方案", 3.5])
        p3 = d / "b.xlsx"
        wb.save(str(p3))
        text, status, _ = extract_text(p3, "xlsx")
        check("xlsx 工作表+行", status == "ok" and "工作表：进度" in text and "写方案" in text, (status, text[:60]))
    except ImportError:
        print(" skip xlsx（缺少 openpyxl）")

    try:
        from pypdf import PdfWriter
        w = PdfWriter()
        w.add_blank_page(width=72, height=72)
        p4 = d / "blank.pdf"
        with open(p4, "wb") as f:
            w.write(f)
        text, status, note = extract_text(p4, "pdf")
        check("空白 pdf → failed(扫描件语义)", status == "failed" and "扫描件" in note, (status, note))
    except ImportError:
        print(" skip pdf（缺少 pypdf）")

    _, status, _ = extract_text(d / "u8.txt", "exe")
    check("unsupported 兜底", status == "unsupported")


def test_api_edges() -> None:
    print("[api] 边界行为")
    with TestClient(app) as client:
        def login(name: str) -> dict:
            r = client.post("/api/v1/auth/dev-login",
                            json={"name": name, "department": "项目部", "position": "工程师"})
            assert r.status_code == 200, r.text
            return {"Authorization": f"Bearer {r.json()['token']}"}

        h1, h2 = login("张三"), login("李四")

        # 413：超过 1MB 上限
        big = b"x" * (1024 * 1024 + 1)
        r = client.post("/api/v1/chat/uploads", headers=h1,
                        files=[("files", ("big.txt", big, "text/plain"))])
        check("超限 413", r.status_code == 413, r.status_code)

        # 415：非法类型
        r = client.post("/api/v1/chat/uploads", headers=h1,
                        files=[("files", ("x.exe", b"MZ", "application/octet-stream"))])
        check("非法类型 415", r.status_code == 415, r.status_code)

        # 空文件 422
        r = client.post("/api/v1/chat/uploads", headers=h1,
                        files=[("files", ("empty.txt", b"", "text/plain"))])
        check("空文件 422", r.status_code == 422, r.status_code)

        # 正常上传（未绑定）
        r = client.post("/api/v1/chat/uploads", headers=h1,
                        files=[("files", ("note.txt", "要点记录".encode("utf-8"), "text/plain"))])
        check("上传成功", r.status_code == 200, r.text)
        att = r.json()[0]

        # 未绑定附件可删除
        r = client.delete(f"/api/v1/chat/uploads/{att['id']}", headers=h1)
        check("未绑定附件可删", r.status_code == 200, r.text)

        # 跨用户访问防护：李四读/删张三的附件、读张三的会话 → 404
        r = client.post("/api/v1/chat/uploads", headers=h1,
                        files=[("files", ("note2.txt", "另一份".encode("utf-8"), "text/plain"))])
        att2 = r.json()[0]
        r = client.get(f"/api/v1/chat/uploads/{att2['id']}", headers=h2)
        check("跨用户读附件 404", r.status_code == 404, r.status_code)
        r = client.delete(f"/api/v1/chat/uploads/{att2['id']}", headers=h2)
        check("跨用户删附件 404", r.status_code == 404, r.status_code)

        r = client.post("/api/v1/chat/sessions", headers=h1, json={})
        sid = r.json()["id"]
        r = client.get(f"/api/v1/chat/sessions/{sid}", headers=h2)
        check("跨用户读会话 404", r.status_code == 404, r.status_code)
        r = client.post(f"/api/v1/chat/sessions/{sid}/messages", headers=h2,
                        json={"content": "越权"})
        check("跨用户发消息 404", r.status_code == 404, r.status_code)

        # 使用他人附件 → 404
        r = client.post(f"/api/v1/chat/sessions/{sid}/messages", headers=h1,
                        json={"content": "hello", "attachment_ids": []})
        check("本人发消息 200", r.status_code == 200, r.text)
        r = client.post("/api/v1/chat/sessions", headers=h2, json={})
        sid2 = r.json()["id"]
        r = client.post(f"/api/v1/chat/sessions/{sid2}/messages", headers=h2,
                        json={"content": "用别人的附件", "attachment_ids": [att2["id"]]})
        check("用他人附件 404", r.status_code == 404, r.status_code)

        # 空消息 422；未知模型 422
        r = client.post(f"/api/v1/chat/sessions/{sid}/messages", headers=h1, json={"content": ""})
        check("空消息 422", r.status_code == 422, r.status_code)
        r = client.post(f"/api/v1/chat/sessions/{sid}/messages", headers=h1,
                        json={"content": "hi", "model": "gpt-9"})
        check("未知模型 422", r.status_code == 422, r.status_code)

        # 多轮：历史应带入（mock 输出不含历史，但消息表应累积 4 条）
        r = client.post(f"/api/v1/chat/sessions/{sid}/messages", headers=h1,
                        json={"content": "第二轮"})
        check("第二轮 200", r.status_code == 200, r.text)
        r = client.get(f"/api/v1/chat/sessions/{sid}", headers=h1)
        check("消息累积 4 条", len(r.json()["messages"]) == 4, len(r.json()["messages"]))

        # 会话不存在的附件 id
        r = client.post(f"/api/v1/chat/sessions/{sid}/messages", headers=h1,
                        json={"content": "hi", "attachment_ids": [999999]})
        check("不存在附件 404", r.status_code == 404, r.status_code)

        # 附件绑定会话后，不能绑定到另一会话
        r = client.post("/api/v1/chat/uploads", headers=h1,
                        files=[("files", ("bind.txt", "绑定测试".encode("utf-8"), "text/plain"))])
        attb = r.json()[0]
        r = client.post(f"/api/v1/chat/sessions/{sid}/messages", headers=h1,
                        json={"content": "带附件", "attachment_ids": [attb["id"]]})
        check("附件绑定成功", r.status_code == 200, r.text)
        r = client.post(f"/api/v1/chat/sessions/{sid2}/messages", headers=h2,
                        json={"content": "x", "attachment_ids": [attb["id"]]})
        check("已绑定附件跨会话/跨用户 404", r.status_code == 404, r.status_code)


def main() -> int:
    test_extract_units()
    test_api_edges()
    if FAILED:
        print(f"\nTEST_CHAT FAILED: {len(FAILED)} 项 -> {FAILED}")
        return 1
    print("\nTEST_CHAT OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
